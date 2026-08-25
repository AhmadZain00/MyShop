from pathlib import Path
import random

from openpyxl import load_workbook

from app.services.excel_service import make_template, HEADERS


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BASE_DIR / "MyShop_Demo_Data.xlsx"


def main():
    # إنشاء القالب الرسمي من نفس البرنامج
    make_template(OUTPUT)

    wb = load_workbook(OUTPUT)
    ws = wb["products"]

    categories = {
        "مشروبات": [
            "مياه معدنية",
            "عصير برتقال",
            "عصير مانجو",
            "مشروب غازي",
            "مشروب طاقة",
        ],
        "أغذية": [
            "أرز",
            "سكر",
            "مكرونة",
            "زيت طعام",
            "تونة",
            "بسكويت",
            "شوكولاتة",
        ],
        "منظفات": [
            "مسحوق غسيل",
            "سائل أطباق",
            "منظف أرضيات",
            "منظف زجاج",
            "مطهر أسطح",
        ],
        "عناية شخصية": [
            "شامبو",
            "صابون",
            "معجون أسنان",
            "فرشاة أسنان",
            "مزيل عرق",
        ],
        "أدوات منزلية": [
            "أكياس قمامة",
            "ورق مطبخ",
            "أكواب بلاستيكية",
            "إسفنجة تنظيف",
        ],
    }

    brands = [
        "النيل",
        "رويال",
        "فاميلي",
        "الصفوة",
        "الواحة",
        "جرين",
        "فريش",
    ]

    units = [
        "قطعة",
        "علبة",
        "عبوة",
        "زجاجة",
        "كيس",
    ]

    random.seed(7)

    product_id = 1001

    for category, products in categories.items():
        for product_name in products:
            for variant in range(1, 3):

                purchase = round(random.uniform(6, 150), 2)
                sale = round(
                    purchase * random.uniform(1.15, 1.35),
                    2,
                )

                state = product_id % 9

                if state == 0:
                    quantity = 0
                    minimum = random.randint(8, 15)
                    note = "نفد المخزون"

                elif state in (1, 2):
                    minimum = random.randint(8, 15)
                    quantity = random.randint(1, minimum)
                    note = "مخزون منخفض"

                else:
                    minimum = random.randint(5, 12)
                    quantity = random.randint(minimum + 5, 120)
                    note = "متوفر"

                ws.append([
                    f"{product_name} {variant}",
                    str(6221000000000 + product_id),
                    f"MY-{product_id:05d}",
                    category,
                    brands[(product_id - 1001) % len(brands)],
                    purchase,
                    sale,
                    quantity,
                    units[(product_id - 1001) % len(units)],
                    minimum,
                    note,
                ])

                product_id += 1

    # حالات واضحة لاختبار Dashboard
    ws.append([
        "منتج تجريبي - مخزون منخفض",
        "6221999000001",
        "DEMO-LOW-01",
        "مشروبات",
        "MyShop Demo",
        20.00,
        29.00,
        3,
        "قطعة",
        15,
        "مخزون منخفض",
    ])

    ws.append([
        "منتج تجريبي - نفد المخزون",
        "6221999000002",
        "DEMO-OUT-01",
        "منظفات",
        "MyShop Demo",
        45.00,
        59.00,
        0,
        "عبوة",
        10,
        "نفد المخزون",
    ])

    # الحفاظ على Barcode و SKU كنص
    for row in ws.iter_rows(
        min_row=2,
        min_col=2,
        max_col=3,
    ):
        row[0].number_format = "@"
        row[1].number_format = "@"

    ws.freeze_panes = "A2"

    wb.save(OUTPUT)

    # تحقق نهائي من أن الأعمدة مطابقة للكود
    check = load_workbook(
        OUTPUT,
        read_only=True,
        data_only=True,
    )

    actual_headers = [
        str(value).strip() if value is not None else ""
        for value in next(
            check["products"].iter_rows(values_only=True)
        )
    ]

    if actual_headers != HEADERS:
        raise RuntimeError(
            "فشل التحقق: عناوين Excel لا تطابق MyShop."
        )

    product_count = check["products"].max_row - 1

    check.close()

    print("تم إنشاء ملف البيانات بنجاح.")
    print(f"المسار: {OUTPUT}")
    print(f"عدد المنتجات: {product_count}")
    print("عناوين الأعمدة مطابقة لكود MyShop.")


if __name__ == "__main__":
    main()