from decimal import Decimal, InvalidOperation
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from app.database import connection as dbconn
from app.database.models import Product, Category, PriceHistory, StockMovement, Sale, User
from sqlalchemy import select
from sqlalchemy.orm import joinedload
from app.services.sales_service import UTC_OFFSET_HOURS

HEADERS = ["اسم المنتج","Barcode","SKU","التصنيف","الماركة","سعر الشراء","سعر البيع","الكمية","الوحدة","الحد الأدنى","ملاحظات"]


def _text(v):
    return str(v).strip() if v is not None else None


def make_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    # Keep identifiers as text so leading zeroes are preserved in Excel.
    for row in ws.iter_rows(min_row=2, max_row=5000, min_col=2, max_col=3):
        row[0].number_format = "@"
        row[1].number_format = "@"
    widths = [28, 18, 18, 20, 18, 16, 16, 14, 12, 14, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("instructions")
    instructions = [
        ["تعليمات MyShop"],
        ["املأ ورقة products ولا تغيّر أسماء الأعمدة."],
        ["Barcode وSKU معرفات نصية؛ لا تضع مسافات قبل/بعد القيمة."],
        ["الأسعار والكمية والحد الأدنى أرقام فقط."],
        ["إذا كان التصنيف غير موجود، ينشئه MyShop تلقائيًا."],
        ["يفضل ضبط Barcode Scanner على إرسال Enter بعد القراءة."],
    ]
    for row in instructions: ins.append(row)
    ins.column_dimensions["A"].width = 100
    wb.save(path)


def export_products(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(HEADERS)
    with dbconn.SessionLocal() as s:
        rows = s.scalars(select(Product).order_by(Product.name)).all()
        for p in rows:
            ws.append([p.name,p.barcode,p.sku,p.category.name if p.category else "",
                       p.brand,float(p.purchase_price),float(p.sale_price),float(p.quantity),
                       p.unit,float(p.min_stock),p.notes])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        row[0].number_format = "@"; row[1].number_format = "@"
    ws.freeze_panes = "A2"
    wb.save(path)


def import_products(path, mode="update", progress=None, cancel_check=None):
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "products" not in wb.sheetnames:
            raise ValueError("ملف Excel لا يحتوي على ورقة products.")
        ws = wb["products"]
        rows = ws.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(rows, ())]
        if header != HEADERS:
            raise ValueError("عناوين الأعمدة غير مطابقة لنموذج MyShop.")
        total = max(0, ws.max_row - 1)
        result = {"added":0,"updated":0,"failed":0,"skipped":0,"errors":[]}
        with dbconn.SessionLocal() as s:
            with s.begin():
                categories = {c.name: c for c in s.scalars(select(Category)).all()}
                for row_number, row in enumerate(rows, start=2):
                    if cancel_check and cancel_check():
                        result["cancelled"] = True
                        break
                    try:
                        # Each row gets a savepoint, so one bad row does not roll back good rows.
                        with s.begin_nested():
                            name = _text(row[0]) or ""
                            if not name:
                                raise ValueError("اسم المنتج فارغ.")
                            barcode = _text(row[1]); sku = _text(row[2])
                            category_name = _text(row[3]) or ""
                            if category_name and category_name not in categories:
                                c = Category(name=category_name); s.add(c); s.flush(); categories[category_name] = c
                            cat = categories.get(category_name)
                            def num(v, label):
                                try: return Decimal(str(v if v is not None and v != "" else 0))
                                except (InvalidOperation, ValueError, TypeError):
                                    raise ValueError(f"{label} غير صالح.")
                            purchase = num(row[5],"سعر الشراء"); sale = num(row[6],"سعر البيع")
                            qty = num(row[7],"الكمية"); minimum = num(row[9],"الحد الأدنى")
                            p = s.scalar(select(Product).where(Product.barcode == barcode)) if barcode else None
                            if p is None and sku:
                                p = s.scalar(select(Product).where(Product.sku == sku))
                            if p:
                                if mode == "ignore":
                                    result["skipped"] += 1
                                    continue
                                old_price = Decimal(p.sale_price); old_qty = Decimal(p.quantity)
                                if mode == "prices":
                                    p.purchase_price, p.sale_price = purchase, sale
                                else:
                                    p.name=name; p.barcode=barcode; p.sku=sku; p.category=cat
                                    p.brand=_text(row[4]); p.purchase_price=purchase; p.sale_price=sale
                                    p.quantity=qty; p.unit=_text(row[8]) or "قطعة"; p.min_stock=minimum; p.notes=_text(row[10])
                                if sale != old_price:
                                    pct = ((sale-old_price)/old_price*100) if old_price else None
                                    s.add(PriceHistory(product_id=p.id, old_price=old_price, new_price=sale,
                                                        change_percent=pct))
                                if mode != "prices" and qty != old_qty:
                                    s.add(StockMovement(product_id=p.id, quantity_before=old_qty,
                                                        quantity_after=qty, delta=qty-old_qty,
                                                        reason="استيراد Excel"))
                                s.flush(); result["updated"] += 1
                            else:
                                p = Product(name=name, barcode=barcode, sku=sku, category=cat,
                                            brand=_text(row[4]), purchase_price=purchase, sale_price=sale,
                                            quantity=qty, unit=_text(row[8]) or "قطعة", min_stock=minimum,
                                            notes=_text(row[10]))
                                s.add(p); s.flush()
                                if qty:
                                    s.add(StockMovement(product_id=p.id, quantity_before=0, quantity_after=qty,
                                                        delta=qty, reason="استيراد Excel"))
                                result["added"] += 1
                    except Exception as exc:
                        result["failed"] += 1
                        result["errors"].append({"row":row_number, "error":str(exc), "data":list(row)})
                    if progress:
                        progress(int((row_number-1) / max(1,total) * 100))
        return result
    finally:
        wb.close()


def export_errors(path, errors):
    wb = Workbook()
    ws = wb.active
    ws.title = "errors"
    ws.append(["رقم الصف","الخطأ"] + HEADERS)
    for e in errors:
        ws.append([e["row"], e["error"]] + e["data"])


SALE_HEADERS = ["التاريخ","الوقت","المنتج","الكمية","سعر الوحدة","الإجمالي","الكاشير"]


def export_sales(path, day=None):
    """Export one local day's sales report to Excel (defaults to today)."""
    day = day or date.today()
    start_local = datetime(day.year, day.month, day.day)
    start = start_local - timedelta(hours=UTC_OFFSET_HOURS)
    end = start + timedelta(days=1)
    wb = Workbook()
    ws = wb.active
    ws.title = "sales"
    ws.append(SALE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    total = Decimal("0")
    with dbconn.SessionLocal() as s:
        rows = s.execute(
            select(Sale, User.username)
            .options(joinedload(Sale.product))
            .outerjoin(User, Sale.sold_by == User.id)
            .where(Sale.sold_at >= start, Sale.sold_at < end)
            .order_by(Sale.sold_at)
        ).all()
        for sale, username in rows:
            local = sale.sold_at + timedelta(hours=UTC_OFFSET_HOURS)
            total += Decimal(sale.total)
            ws.append([local.strftime("%Y-%m-%d"), local.strftime("%H:%M:%S"),
                       sale.product.name, float(sale.quantity), float(sale.unit_price),
                       float(sale.total), username or "-"])
    ws.append(["", "", "", "", "الإجمالي", float(total), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i, width in enumerate([14, 12, 30, 12, 14, 14, 16], 1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return {"count": len(rows), "total": total}
    wb.save(path)